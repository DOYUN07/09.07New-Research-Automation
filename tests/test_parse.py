"""파서 검증 — 네트워크 없이 저장된 마크업으로 확인한다.

실행: python -m tests.test_parse
"""
from __future__ import annotations

import sys
from datetime import date

from src.config import Institution, load_config
from src.filters import FilterStats, apply_filters
from src.parse import parse
from tests import fixtures

TODAY = date(2026, 9, 7)
FAILS: list[str] = []


def check(label: str, got, want) -> None:
    if got != want:
        FAILS.append(f"{label}\n    기대: {want!r}\n    실제: {got!r}")
        print(f"  ✗ {label}: {got!r} != {want!r}")
    else:
        print(f"  ✓ {label}")


def inst(**kw) -> Institution:
    kw.setdefault("id", "t")
    kw.setdefault("name", "테스트")
    kw.setdefault("url", "https://example.com/list")
    kw.setdefault("base", "https://example.com")
    return Institution(**kw)


def P(html: str, institution: Institution):
    """기준일을 2026-09-07로 고정해서 파싱한다 (미래 날짜 판정이 결과에 영향을 주므로)."""
    return parse(html, institution, TODAY)


def main() -> int:
    print("\n[1] 부산테크노파크형 — 접수기간과 게시일이 따로 있는 표")
    r = P(fixtures.BTP, inst(id="btp", base="https://www.btp.or.kr/kor/CMS/Board/Board.do"))
    check("행 수", len(r), 4)
    check("제목", r[0].title, "Age-Tech 종합지원센터 운영 사업 TRL기반 기술성장 맞춤 지원 공고")
    check("게시일", r[0].posted, date(2026, 9, 7))
    check("마감일", r[0].deadline, date(2026, 9, 30))
    check("링크 절대경로", r[0].url.startswith("https://www.btp.or.kr/"), True)
    check("마감 상태 인식", r[1].closed_flag, True)

    print("\n[2] 광주테크노파크형 — 기간만 있고 게시일 컬럼 없음, 번호가 th")
    r = P(fixtures.GJTP, inst(id="gjtp", base="https://www.gjtp.or.kr/home/business.cs"))
    check("행 수", len(r), 3)
    check("게시일=기간 시작", r[0].posted, date(2026, 9, 3))
    check("마감일=기간 끝", r[0].deadline, date(2026, 9, 14))

    print("\n[3] 농림축산식품부형 — 날짜가 dd.date 안")
    r = P(fixtures.MAFRA, inst(id="mafra", base="https://www.mafra.go.kr"))
    check("행 수", len(r), 3)
    check("제목에서 '새글' 제거", r[0].title, "스마트농업 클라우드 실증 지원사업 공고")
    check("게시일", r[0].posted, date(2026, 9, 4))

    print("\n[4] IRIS형 — ul/li 구조 + onclick 링크")
    r = P(
        fixtures.IRIS,
        inst(
            id="iris",
            base="https://www.iris.go.kr",
            row_selector="ul.dbody > li",
            link_from_onclick=r"_view\('(\d+)'",
            detail_url="https://www.iris.go.kr/contents/retrieveBsnsAncmView.do?ancmId={id}&ancmPrg=ancmIng",
        ),
    )
    check("행 수", len(r), 3)
    check("게시일", r[0].posted, date(2026, 9, 7))
    check(
        "onclick에서 상세 URL 조립",
        r[0].url,
        "https://www.iris.go.kr/contents/retrieveBsnsAncmView.do?ancmId=023977&ancmPrg=ancmIng",
    )

    print("\n[4-b] IRIS형 — 선택자 없이 자동 인식되는지")
    r_auto = P(
        fixtures.IRIS,
        inst(
            id="iris",
            base="https://www.iris.go.kr",
            link_from_onclick=r"_view\('(\d+)'",
            detail_url="https://www.iris.go.kr/x?ancmId={id}",
        ),
    )
    check("자동 인식 행 수", len(r_auto), 3)

    print("\n[5] 중소벤처기업부형 — onclick doBbsFView + 중첩 div 안의 신청기간")
    r = P(
        fixtures.MSS,
        inst(
            id="mss",
            base="https://www.mss.go.kr",
            link_from_onclick=r"doBbsFView\('\d+','(\d+)'",
            detail_url="https://www.mss.go.kr/site/smba/ex/bbs/View.do?cbIdx=310&bcIdx={id}&parentSeq={id}",
        ),
    )
    check("행 수", len(r), 3)
    check("게시일", r[0].posted, date(2026, 9, 7))
    check("마감일=신청기간 종료", r[0].deadline, date(2026, 10, 6))
    check("상세 URL bcIdx", "bcIdx=1071012" in r[0].url, True)

    print("\n[6] 부산사회서비스원형 — 2자리 연도(26.09.05)")
    r = P(fixtures.BUSAN_PASS, inst(id="bp", base="https://busan.pass.or.kr/SW_bbs/notice/"))
    check("행 수", len(r), 3)
    check("2자리 연도 해석", r[0].posted, date(2026, 9, 5))

    print("\n[7] K-Startup형 — 카드형 목록 + go_view(id)")
    r = P(
        fixtures.KSTARTUP,
        inst(
            id="ks",
            base="https://www.k-startup.go.kr",
            link_from_onclick=r"go_view\((\d+)\)",
            detail_url="https://www.k-startup.go.kr/web/contents/bizpbanc-ongoing.do?schM=view&pbancSn={id}&page=1",
        ),
    )
    check("행 수", len(r), 3)
    check("게시일", r[0].posted, date(2026, 9, 4))
    check("마감일", r[0].deadline, date(2026, 9, 15))
    check("상세 URL", "pbancSn=179111" in r[0].url, True)

    print("\n[8] 게시일이 아예 없는 게시판")
    r = P(fixtures.NO_DATE, inst(id="nd", base="https://www.khidi.or.kr"))
    check("행 수", len(r), 3)
    check("게시일 None", r[0].posted, None)
    check("제목", r[0].title, "[공고] 2026년 1차 고령친화우수제품 지정 공고")

    # -------------------------------- 1차 실전 진단(2026-09-07)에서 발견된 문제들
    print("\n[A] 첨부파일 링크를 제목으로 착각하지 않는지 (산업통상부에서 발생)")
    r = P(fixtures.ATTACH_TRAP, inst(id="motir", base="https://www.motir.go.kr"))
    check("행 수", len(r), 3)
    check("제목이 파일명이 아님", r[0].title, "2026년도 산업통상부-에너지공기업 기술나눔")
    check("'.hwpx' 미포함", ".hwpx" in r[0].title, False)
    check("'다운로드' 미포함", "다운로드" in r[0].title, False)
    check("링크가 첨부파일이 아님", "download" in r[0].url, False)
    check("링크가 본문", "/71310/view" in r[0].url, True)

    print("\n[B] 제목이 두 번 이어붙는 현상 (부산테크노파크에서 발생)")
    r = P(fixtures.DOUBLED_TITLE, inst(id="dbl", base="https://example.com"))
    check("행 수", len(r), 3)
    check("제목 1회만", r[0].title, "신중년 디지털 전환 지원사업 공고")
    check("제목 1회만 (2)", r[1].title, "고령친화 서비스 실증 참여기업 모집")

    print("\n[C] 한 행에 기간이 두 개일 때 게시일이 미래로 잡히지 않는지 (연구개발특구진흥재단)")
    r = P(fixtures.TWO_PERIODS, inst(id="innopolis", base="https://pms.innopolis.or.kr"))
    check("행 수", len(r), 3)
    check("게시일=공고기간 시작", r[0].posted, date(2026, 8, 21))
    check("게시일이 오늘 이전", r[0].posted <= TODAY, True)
    check("마감일=가장 늦은 종료일", r[0].deadline, date(2026, 9, 21))
    check("2행 게시일", r[1].posted, date(2026, 9, 1))
    check("2행 마감일", r[1].deadline, date(2026, 10, 5))

    # -------------------------------- 2차 실전 진단(2026-09-07)에서 발견된 문제들
    print("\n[D] 목록에 마감일시만 있을 때 (전북테크노파크)")
    r = P(fixtures.DEADLINE_ONLY, inst(id="jbtp", base="https://jbcis.jbtp.or.kr"))
    check("행 수", len(r), 3)
    check("마감일시를 게시일로 쓰지 않음", r[0].posted, None)
    check("마감일로 인식", r[0].deadline, date(2026, 12, 31))
    check("3행 마감일", r[2].deadline, date(2026, 9, 30))

    print("\n[E] 두 span 내용이 완전히 같지 않은 중복 제목 (부산테크노파크)")
    r = P(fixtures.PARTIAL_DOUBLE, inst(id="btp2", base="https://www.btp.or.kr"))
    check("행 수", len(r), 3)
    check(
        "긴 쪽만 채택",
        r[0].title,
        "기업성장기반 글로벌 하이메디 허브 특구 상생협력사업 기업지원모집 공고(4차) 재공고",
    )
    check("말줄임표 미포함", "..." in r[0].title, False)
    check("2행도 1회만", r[1].title.count("글로벌시장"), 1)
    check("완전 동일한 경우도 1회", r[2].title, "시니어 돌봄로봇 실증 참여기업 모집 공고")

    print("\n[F] 게시판이 <header> 안에 있어도 인식 (NIPA 0건 회귀 방지)")
    r = P(fixtures.INSIDE_HEADER, inst(id="nipa", base="https://www.nipa.kr"))
    check("행 수", len(r), 3)
    check("제목", r[0].title, "2026년 아태 AI 특화지구(AHAP) 조성 사업 공고")
    check("class에 down이 있어도 링크 유지", "/home/2-2/16921" in r[0].url, True)
    check("게시일", r[0].posted, date(2026, 9, 3))

    # -------------------------------- 3차 실전 진단(2026-09-07)에서 발견된 문제들
    print("\n[G] 행 안에 첨부파일 <li>가 있어도 인식 (중소벤처기업부)")
    r = P(
        fixtures.ROW_WITH_LI,
        inst(
            id="mss",
            base="https://www.mss.go.kr",
            link_from_onclick=r"doBbsFView\('\d+','(\d+)'",
            detail_url="https://www.mss.go.kr/site/smba/ex/bbs/View.do?cbIdx=310&bcIdx={id}&parentSeq={id}",
        ),
    )
    check("행 수", len(r), 3)
    check(
        "제목에 메타데이터가 안 붙음",
        r[0].title,
        "2026년 중소기업 스마트서비스 지원사업 참여기업 모집 공고(A/S지원)",
    )
    check("'담당부서' 미포함", "담당부서" in r[0].title, False)
    check("게시일", r[0].posted, date(2026, 9, 7))
    check("마감일=신청기간 종료", r[0].deadline, date(2026, 10, 6))
    check("상세 URL", "bcIdx=1071012" in r[0].url, True)

    print("\n[H] 행마다 감싸는 div가 따로 있어 형제가 아닌 목록 (부산시민운동지원센터)")
    r = P(fixtures.WRAPPED_ROWS, inst(id="ng", base="https://www.ngocenter.or.kr/business/"))
    check("행 수", len(r), 3)
    check("제목", r[0].title, "[활동가커뮤니티지원사업] 든든 커뮤니티 큰모임 (9/12)")
    check("게시일", r[0].posted, date(2026, 9, 2))
    check("상대경로 결합", r[0].url, "https://www.ngocenter.or.kr/business/view?scti=0&no=2904")
    check("헤더 행 제외", any("제목" == n.title for n in r), False)

    # ---------------------------------------------------------------- 필터
    print("\n[9] 필터 규칙")
    cfg = load_config()
    cfg.include_keywords = ["Age-Tech", "지원사업", "고령", "AX", "실증", "돌봄", "클라우드"]
    cfg.exclude_keywords = ["채용", "포상", "구제역", "퇴직공무원"]
    cfg.lookback_days = 10
    cfg.max_per_institution = 10

    parsed = P(fixtures.BTP, inst(id="btp", base="https://www.btp.or.kr/"))
    st = FilterStats()
    kept = apply_filters(parsed, cfg, TODAY, set(), st)
    titles = [n.title for n in kept]
    check("채용 공고 제외", any("채용" in t for t in titles), False)
    check("마감된 공고 제외", any("커피산업" in t for t in titles), False)
    check("Age-Tech 공고 채택", any("Age-Tech" in t for t in titles), True)

    print("\n[9-b] 게시일 10일 초과 공고 제외")
    parsed = P(fixtures.BUSAN_PASS, inst(id="bp", base="https://busan.pass.or.kr/"))
    cfg.include_keywords = ["사회서비스", "돌봄"]
    cfg.exclude_keywords = []
    st = FilterStats()
    kept = apply_filters(parsed, cfg, TODAY, set(), st)
    check("07.29 공고 제외됨", any("역량강화" in n.title for n in kept), False)
    check("09.05 / 09.02 공고 채택", len(kept), 2)

    print("\n[9-c] 중복 발송 방지")
    st = FilterStats()
    already = {kept[0].key}
    kept2 = apply_filters(
        P(fixtures.BUSAN_PASS, inst(id="bp", base="https://busan.pass.or.kr/")),
        cfg,
        TODAY,
        already,
        st,
    )
    check("이미 보낸 공고 제외", len(kept2), 1)
    check("중복 카운트", st.dropped_duplicate, 1)

    print("\n[9-d] 마감일 미상 공고는 '확인필요'로 포함")
    cfg.include_keywords = ["고령", "에이지테크", "지정"]
    st = FilterStats()
    kept3 = apply_filters(
        P(fixtures.NO_DATE, inst(id="nd", base="https://www.khidi.or.kr")),
        cfg,
        TODAY,
        set(),
        st,
    )
    check("마감일 미상도 채택됨", len(kept3) >= 2, True)
    check("마감일 None", kept3[0].deadline, None)

    print("\n[9-e] unknown_deadline: exclude 로 바꾸면 제외")
    cfg.unknown_deadline = "exclude"
    st = FilterStats()
    kept4 = apply_filters(
        P(fixtures.NO_DATE, inst(id="nd", base="https://www.khidi.or.kr")),
        cfg,
        TODAY,
        set(),
        st,
    )
    check("전부 제외", len(kept4), 0)

    # ---------------------------------------------------------------- 누적 기록
    print("\n[10] 누적 기록 (엑셀)")
    import tempfile
    from pathlib import Path

    from src import archive

    src_notices = P(fixtures.BTP, inst(id="btp", name="부산테크노파크", base="https://www.btp.or.kr/"))
    for n in src_notices:
        n.matched_keywords = ["테스트"]

    with tempfile.TemporaryDirectory() as td:
        csvp = Path(td) / "archive.csv"
        xlsxp = Path(td) / "공고누적.xlsx"

        rows, added = archive.add([], src_notices, date(2026, 9, 7))
        check("1일차 추가 건수", added, len(src_notices))
        archive.save_csv(rows, csvp)
        check("CSV 저장됨", csvp.exists(), True)

        # 같은 공고를 다시 넣어도 늘지 않아야 한다
        rows2 = archive.load_rows(csvp)
        check("CSV 다시 읽기", len(rows2), len(src_notices))
        rows2, added2 = archive.add(rows2, src_notices, date(2026, 9, 8))
        check("중복은 추가 안 됨", added2, 0)
        check("누적 건수 유지", len(rows2), len(src_notices))

        made = archive.build_xlsx(rows2, xlsxp)
        if made is None:
            print("  (openpyxl 미설치 — 엑셀 검증 건너뜀)")
        else:
            from openpyxl import load_workbook

            ws = load_workbook(xlsxp).active
            check("시트 이름", ws.title, "공고누적")
            check("데이터 행 수", ws.max_row - 1, len(src_notices))
            check(
                "헤더",
                [c.value for c in ws[1]],
                ["발견일", "기관", "공고명", "게시일", "마감일", "매칭 키워드", "링크"],
            )
            check("'키' 컬럼은 제외", "키" in [c.value for c in ws[1]], False)
            check("고정틀", ws.freeze_panes, "A2")
            painted = sum(
                1
                for row in ws.iter_rows()
                for c in row
                if c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None, "00000000")
            )
            check("색 채운 셀 없음", painted, 0)
            check("링크는 하이퍼링크", ws.cell(2, 7).value, "바로가기")

    print()
    if FAILS:
        print(f"실패 {len(FAILS)}건")
        for f in FAILS:
            print("  - " + f)
        return 1
    print("모든 검증 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
